"""Unit tests for app.utils.document_text.extract_text.

Covers the bug this module fixes: PDFs uploaded as ``transcript`` were
previously UTF-8-decoded as raw bytes (garbage), and PDFs uploaded as
``supporting_document`` never had any text extracted at all.
"""
from __future__ import annotations

import io

import pytest
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from app.utils.document_text import DocumentTextExtractionError, extract_text


def _build_pdf_bytes(text: str) -> bytes:
    """Builds a minimal, real, one-page PDF containing ``text`` as content."""
    writer = PdfWriter()
    page = writer.add_blank_page(width=200, height=200)

    font = DictionaryObject(
        {
            NameObject("/Type"): NameObject("/Font"),
            NameObject("/Subtype"): NameObject("/Type1"),
            NameObject("/BaseFont"): NameObject("/Helvetica"),
        }
    )
    font_ref = writer._add_object(font)
    page[NameObject("/Resources")] = DictionaryObject(
        {NameObject("/Font"): DictionaryObject({NameObject("/F1"): font_ref})}
    )

    content = DecodedStreamObject()
    content.set_data(f"BT /F1 12 Tf 10 100 Td ({text}) Tj ET".encode("latin-1"))
    content_ref = writer._add_object(content)
    page[NameObject("/Contents")] = content_ref

    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_blank_pdf_bytes() -> bytes:
    """A structurally valid PDF with a page but no text content at all."""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _build_docx_bytes() -> bytes:
    document = Document()
    document.add_heading("Apache modernization", level=1)
    document.add_paragraph("Move the workload to an Azure-native platform.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Owner"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "Operations"
    table.cell(1, 1).text = "Zero-downtime migration"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _build_blank_docx_bytes() -> bytes:
    document = Document()
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_plain_text_transcript_is_decoded_as_utf8() -> None:
    text = extract_text(
        content=b"hello world",
        content_type="text/plain",
        file_name="transcript.txt",
    )
    assert text == "hello world"


def test_markdown_transcript_is_decoded_as_plain_text() -> None:
    text = extract_text(
        content=b"# Requirements\n\n- Must support SSO\n- Must log audit events\n",
        content_type="text/markdown",
        file_name="requirements.md",
    )
    assert text == "# Requirements\n\n- Must support SSO\n- Must log audit events\n"


def test_markdown_transcript_with_generic_content_type_is_decoded_as_plain_text() -> None:
    """Browsers commonly send a generic/empty content type for .md files -
    extraction must not depend on the browser correctly tagging it."""
    text = extract_text(
        content=b"## Call Transcript\n\nCustomer: We need real-time reporting.\n",
        content_type="application/octet-stream",
        file_name="call-notes.md",
    )
    assert text == "## Call Transcript\n\nCustomer: We need real-time reporting.\n"


def test_pdf_detected_by_content_type_is_parsed_for_text() -> None:
    pdf_bytes = _build_pdf_bytes("Hello PDF")
    text = extract_text(content=pdf_bytes, content_type="application/pdf", file_name="doc")
    assert "Hello PDF" in text


def test_pdf_detected_by_file_extension_is_parsed_for_text() -> None:
    pdf_bytes = _build_pdf_bytes("Extracted via extension")
    text = extract_text(
        content=pdf_bytes, content_type="application/octet-stream", file_name="requirements.pdf"
    )
    assert "Extracted via extension" in text


def test_docx_detected_by_content_type_extracts_paragraphs_and_tables() -> None:
    text = extract_text(
        content=_build_docx_bytes(),
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        file_name="document",
    )

    assert "Apache modernization" in text
    assert "Move the workload to an Azure-native platform." in text
    assert "Owner | Requirement" in text
    assert "Operations | Zero-downtime migration" in text


def test_docx_detected_by_file_extension_is_parsed_for_text() -> None:
    text = extract_text(
        content=_build_docx_bytes(),
        content_type="application/octet-stream",
        file_name="requirements.docx",
    )

    assert "Apache modernization" in text


def test_blank_pdf_with_no_text_raises_extraction_error() -> None:
    pdf_bytes = _build_blank_pdf_bytes()
    with pytest.raises(DocumentTextExtractionError):
        extract_text(content=pdf_bytes, content_type="application/pdf", file_name="blank.pdf")


def test_malformed_pdf_bytes_raise_extraction_error() -> None:
    with pytest.raises(DocumentTextExtractionError):
        extract_text(content=b"not a real pdf", content_type="application/pdf", file_name="bad.pdf")


def test_blank_docx_with_no_text_raises_extraction_error() -> None:
    with pytest.raises(DocumentTextExtractionError):
        extract_text(
            content=_build_blank_docx_bytes(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_name="blank.docx",
        )


def test_malformed_docx_bytes_raise_extraction_error() -> None:
    with pytest.raises(DocumentTextExtractionError):
        extract_text(
            content=b"not a real docx",
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            file_name="bad.docx",
        )


def test_unknown_binary_type_is_not_decoded_as_customer_evidence() -> None:
    with pytest.raises(DocumentTextExtractionError, match="Unsupported local document type"):
        extract_text(
            content=b"PK\x03\x04binary",
            content_type="application/octet-stream",
            file_name="evidence.zip",
        )


def test_invalid_utf8_is_not_replaced_with_lossy_characters() -> None:
    with pytest.raises(DocumentTextExtractionError, match="UTF-8"):
        extract_text(
            content=b"customer requirement: \xff",
            content_type="text/plain",
            file_name="requirements.txt",
        )
